import io
import unittest

import openpyxl

from calculator import calculate, export_workbook, inspect


def fixture(*, demand=10, available=100, domestic=200, shipped=0, status='已完成',
            factory=0, pack=20, duplicate=False, missing_inventory=False):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    def sheet(name, heads, rows):
        s = wb.create_sheet(name)
        s.append(heads)
        for row in rows:
            s.append(row)
        return s
    sheet('货件补货计算-创货件计划单', ['FNSKU','产品名称','单箱数量','预期日销'],
          [['X1','=HYPERLINK("bad")',pack,12]])
    heads = ['店铺','FNSKU','FBA可售','FBA待调仓','FBA调仓中','FBA标发在途',
             'FBA实际在途','FBA入库中','AWD可用量','AWD标发在途','AWD实际在途',
             'AWD补货至FBA标发在途','AWD补货至FBA实际在途']
    row = ['Shop','X1',available,0,0,50,40,0,30,0,0,40,20]
    inv = sheet('FBA库存明细今天', heads, [] if missing_inventory else [row])
    if duplicate:
        inv.append(row)
    sheet('listing报告今天',['店铺','FNSKU','SKU','7日均销量','30日均销量','昨日销量'],
          [['Shop','X1','SKU1',demand,8,9]])
    sheet('计划单今天',['店铺','条码(FNSKU)','总数量','发货数量','计划状态'],
          [['本地店铺','X1',domestic,shipped,status]])
    sheet('厂家未出今天',['条码(FNSKU)','总数量'],[['X1',factory]])
    ship = sheet('已创数据今天',['条码(FNSKU或GTIN)','单箱数量','箱数','总数量'],
                 [['X1',20,2,39]])
    # A formatting-only last row must not expand the read to one million rows.
    ship.cell(1048576, 1).number_format = '0'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


CONFIRMED = {'plan_store':'本地店铺','coverage_mode':'manual','data_confirmed':True,
             'production_days':2,'shipping_days':3,'target_days':10,'safety_days':1}


class CalculatorTests(unittest.TestCase):
    def row(self, data=None, options=None):
        return calculate(data or fixture(), options or CONFIRMED)['results'][0]

    def test_fba_and_procurement_are_separate_and_supply_is_not_duplicated(self):
        # FBA: 10*(3+10+1)-100-50-30 < 0. Procurement: 160-180-200 < 0.
        r = self.row()
        self.assertEqual(r['FBA在途估计'], 50)
        self.assertEqual(r['建议FBA发货量'], 0)
        self.assertEqual(r['建议采购量'], 0)

    def test_cartons_and_domestic_limit(self):
        r = self.row(fixture(demand=30, available=0, domestic=75))
        # Demand 420 minus transit 50 and AWD 30 = 340; only 3 domestic cartons.
        self.assertEqual(r['FBA补货需求量'], 340)
        self.assertEqual(r['建议FBA发货量'], 60)
        self.assertEqual(r['国内现货不足量'], 280)
        # Purchase: 480-50-30-75=325 -> 340 units.
        self.assertEqual(r['建议采购量'], 340)

    def test_current_completed_balance_is_not_deducted_twice(self):
        r = self.row(fixture(demand=30, available=0, domestic=75, shipped=40))
        self.assertEqual(r['国内仓候选结余'], 75)
        self.assertEqual(r['建议FBA发货量'], 60)

    def test_unconfirmed_does_not_issue_orders(self):
        r = self.row(options={'plan_store':'本地店铺'})
        self.assertIsNone(r['建议FBA发货量'])
        self.assertIsNone(r['建议采购量'])
        self.assertIn('FBA补货需求量', r)

    def test_bag_factory_quantities_are_never_product_supply(self):
        content = fixture(demand=30, available=0, domestic=0, factory=200)
        a = self.row(content)
        b = self.row(fixture(demand=30, available=0, domestic=0, factory=0))
        self.assertEqual(a['建议采购量'],b['建议采购量'])
        self.assertEqual(a['已有采购候选未交付'],0)

    def test_coverage_tiers_match_original_thresholds(self):
        for demand, days in [(19,80),(19.1,90),(49,90),(49.1,100)]:
            with self.subTest(demand=demand):
                r=self.row(fixture(demand=demand),{**CONFIRMED,'coverage_mode':'tiered'})
                self.assertEqual(r['采用覆盖天数'],days)

    def test_confirmed_lead_time_defaults(self):
        report=calculate(fixture(),{'plan_store':'本地店铺'})
        self.assertEqual(report['settings']['production_days'],45)
        self.assertEqual(report['settings']['shipping_days'],50)

    def test_prior_shipments_deduct_from_open_orders(self):
        r=self.row(fixture(domestic=200,shipped=40,status='已采购'))
        self.assertEqual(r['已有采购候选未交付'],160)

    def test_unknown_plan_status_blocks_final_quantities(self):
        r = self.row(fixture(status='缺料'))
        self.assertIsNone(r['建议采购量'])
        self.assertEqual(r['状态'], '待确认')

    def test_orders_are_not_domestic_ready_stock(self):
        r = self.row(fixture(demand=30, available=0, status='已采购'))
        self.assertEqual(r['建议FBA发货量'],0)
        self.assertEqual(r['已有采购候选未交付'],200)

    def test_duplicate_source_key_blocks_product(self):
        r = self.row(fixture(duplicate=True))
        self.assertEqual(r['状态'],'数据错误')
        self.assertIsNone(r['建议采购量'])

    def test_missing_record_is_not_zero_stock(self):
        self.assertEqual(self.row(fixture(missing_inventory=True))['状态'],'数据错误')

    def test_missing_sales_is_not_zero(self):
        self.assertEqual(self.row(fixture(demand=None))['状态'],'数据错误')

    def test_zero_sales_requires_review(self):
        r = self.row(fixture(demand=0))
        self.assertEqual(r['状态'],'零销量待复核')
        self.assertIsNone(r['建议采购量'])

    def test_invalid_numbers_and_packs(self):
        for field in [dict(pack=0),dict(pack=-1),dict(pack=2.5),dict(available=-1)]:
            with self.subTest(field=field):
                self.assertEqual(self.row(fixture(**field))['状态'],'数据错误')

    def test_no_cross_store_plan(self):
        r = self.row(options={**CONFIRMED,'plan_store':'Wrong'})
        self.assertIsNone(r['建议采购量'])

    def test_formatted_millionth_row_does_not_expand_records(self):
        info = inspect(fixture())
        s = next(x for x in info['sheets'] if x['name'].startswith('已创数据'))
        self.assertEqual(s['rows'],1)

    def test_export_is_readable_and_formula_safe(self):
        report = calculate(fixture(), CONFIRMED)
        report['results'][0]['产品名称'] = '=HYPERLINK("bad")'
        b = export_workbook(report)
        w = openpyxl.load_workbook(io.BytesIO(b))
        self.assertEqual(len(w.sheetnames),6)
        cell = w['补货与采购']['C2']
        self.assertNotEqual(cell.data_type,'f')
        self.assertTrue(cell.value.startswith("'="))

    def test_shipment_carton_check_does_not_invent_received_units(self):
        r = calculate(fixture(), CONFIRMED)['shipments'][0]
        self.assertEqual(r['装箱差异'],-1)
        self.assertIsNone(r['接收数量'])
        self.assertIsNone(r['差异'])

    def test_bad_file_and_bad_parameters(self):
        with self.assertRaises(ValueError):
            inspect(b'not excel')
        with self.assertRaises(ValueError):
            self.row(options={**CONFIRMED,'target_days':0})


if __name__ == '__main__':
    unittest.main()
