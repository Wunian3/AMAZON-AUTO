"""Workbook adapters and auditable, unit-based replenishment calculations."""
from __future__ import annotations

import io
import math
import zipfile
from collections import defaultdict
from datetime import date, datetime
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.styles import Font, PatternFill

MAX_UPLOAD = 20 * 1024 * 1024
ROLES = {
    'catalog': ('货件补货计算',), 'inventory': ('FBA库存明细',),
    'sales': ('listing报告',), 'plan': ('计划单',),
    'factory_sent': ('厂家已出',), 'factory_pending': ('厂家未出',),
    'awd': ('ADW库存', 'AWD库存'), 'shipments': ('已创数据',),
}
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def text(value):
    return '' if value is None else str(value).strip()


def number(value, label, *, optional=False):
    if value in (None, '') and optional:
        return 0.0
    if isinstance(value, bool):
        raise ValueError(f'{label}不是有效数量')
    try:
        n = float(value)
    except (ValueError, TypeError):
        raise ValueError(f'{label}缺失或不是有效数字：{value}') from None
    if not math.isfinite(n) or n < 0:
        raise ValueError(f'{label}必须是非负有限数字')
    return n


def read_workbook(content):
    """Bound actual XML rows instead of trusting Excel's formatted dimensions."""
    if len(content) > MAX_UPLOAD:
        raise ValueError('文件超过 20 MB')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if sum(x.file_size for x in infos) > 100 * 1024 * 1024 or len(infos) > 2000:
                raise ValueError('工作簿解压后过大')
            bounds = {}
            for name in archive.namelist():
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    root = ET.fromstring(archive.read(name))
                    actual = [int(r.attrib['r']) for r in root.findall('s:sheetData/s:row', NS)
                              if any(c.find('s:v', NS) is not None or c.find('s:is', NS) is not None
                                     or c.find('s:f', NS) is not None for c in r)]
                    bounds['/' + name] = max(actual, default=1)
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        tables = {}
        for sheet in wb:
            limit = bounds.get(sheet._worksheet_path, bounds.get('/' + sheet._worksheet_path))
            if limit is None or limit > 30000 or sheet.max_column > 200:
                raise ValueError(f'工作表 {sheet.title} 超出支持范围（3万行、200列）')
            rows = sheet.iter_rows(max_row=limit, values_only=True)
            headers = [text(x) for x in next(rows)]
            # Repeated result headings are irrelevant; raw adapters use unique source headers.
            tables[sheet.title] = [dict(zip(headers, row), __row__=i)
                                   for i, row in enumerate(rows, 2) if any(x is not None for x in row)]
        wb.close()
        return tables
    except (zipfile.BadZipFile, ET.ParseError, KeyError, OSError) as exc:
        raise ValueError('无法读取 Excel，请上传未加密的 .xlsx 文件') from exc


def roles_for(tables):
    roles = {}
    for role, prefixes in ROLES.items():
        matches = [s for s in tables if any(s.startswith(p) for p in prefixes)]
        if len(matches) > 1:
            raise ValueError(f'{role} 存在多个版本：{"、".join(matches)}。请保留本次要核算的一份。')
        if matches:
            roles[role] = matches[0]
    for role in ('catalog', 'inventory', 'sales'):
        if role not in roles:
            raise ValueError(f'缺少 {ROLES[role][0]} 工作表。第一版支持与你提供的样表相同的表头。')
    return roles


def inspect(content):
    tables = read_workbook(content)
    roles = roles_for(tables)
    shops = sorted({text(r.get('店铺')) for r in tables[roles['sales']] if text(r.get('店铺'))})
    return {'sheets': [{'name': k, 'rows': len(v)} for k, v in tables.items()],
            'shops': shops, 'roles': roles}


def unique_index(rows, shop, label):
    index, duplicates = {}, set()
    for row in rows:
        if shop and text(row.get('店铺')) != shop:
            continue
        key = text(row.get('FNSKU'))
        if not key:
            continue
        if key in index:
            duplicates.add(key)
        index[key] = row
    return index, duplicates


def settings(options):
    cfg = {k: number(options.get(k, default), k) for k, default in
           [('production_days', 45), ('shipping_days', 50), ('target_days', 80), ('safety_days', 7)]}
    if cfg['target_days'] <= 0 or any(v > 730 for v in cfg.values()):
        raise ValueError('备货天数须大于0，所有天数不得超过730')
    cfg['sales_basis'] = options.get('sales_basis', '7日均销量')
    if cfg['sales_basis'] not in ('7日均销量', '30日均销量', '昨日销量', '预期日销'):
        raise ValueError('不支持的销量口径')
    cfg['coverage_mode'] = options.get('coverage_mode', 'tiered')
    if cfg['coverage_mode'] not in ('tiered', 'manual'):
        raise ValueError('请选择销量分档或手动覆盖天数')
    cfg['plan_store'] = text(options.get('plan_store'))
    cfg['as_of'] = date.fromisoformat(options.get('as_of', date.today().isoformat()))
    cfg['data_confirmed'] = options.get('data_confirmed') is True
    return cfg


def calculate(content, options):
    tables = read_workbook(content)
    roles = roles_for(tables)
    cfg = settings(options)
    sales_rows = tables[roles['sales']]
    shops = sorted({text(r.get('店铺')) for r in sales_rows if text(r.get('店铺'))})
    shop = text(options.get('shop'))
    if not shop and len(shops) == 1:
        shop = shops[0]
    if shop not in shops:
        raise ValueError('请选择报表中的店铺，不能混合不同店铺核算')
    inv, inv_dup = unique_index(tables[roles['inventory']], shop, '库存')
    sales, sales_dup = unique_index(sales_rows, shop, '销量')
    issues, results, source_rows = [], [], []

    def issue(key, message):
        issues.append({'FNSKU': key, '问题': message})

    plans = defaultdict(list)
    # Confirmed business rule: factory sheets describe bags, not sellable product supply.
    for role in ('plan',):
        for r in tables.get(roles.get(role), []):
            key = text(r.get('条码(FNSKU)'))
            if key:
                plans[key].append((role, r))
    required_inv = ['FBA可售', 'FBA待调仓', 'FBA调仓中', 'FBA标发在途',
                    'FBA实际在途', 'FBA入库中', 'AWD可用量', 'AWD标发在途',
                    'AWD实际在途', 'AWD补货至FBA标发在途', 'AWD补货至FBA实际在途']
    catalogs = tables[roles['catalog']]
    catalog_counts = defaultdict(int)
    for r in catalogs:
        if text(r.get('FNSKU')):
            catalog_counts[text(r['FNSKU'])] += 1

    for catalog in catalogs:
        key = text(catalog.get('FNSKU'))
        if not key:
            continue
        out = {'FNSKU': key, 'SKU': '', '产品名称': text(catalog.get('产品名称')),
               '状态': '待确认', '建议FBA发货量': None, '建议采购量': None}
        results.append(out)
        try:
            if catalog_counts[key] > 1:
                raise ValueError('商品主表 FNSKU 重复，请去重后核算')
            if key in inv_dup or key in sales_dup:
                raise ValueError('库存或销量报表的店铺+FNSKU重复，已停止该商品计算')
            if key not in inv or key not in sales:
                raise ValueError('缺少匹配的库存或销量记录，不能当作零库存')
            inventory, sale = inv[key], sales[key]
            out['SKU'] = text(sale.get('SKU') or sale.get('MSKU'))
            # Never silently reuse cached formula results from the old summary sheet.
            demand = number(catalog.get('预期日销') if cfg['sales_basis'] == '预期日销'
                            else sale.get(cfg['sales_basis']), cfg['sales_basis'])
            pack = number(catalog.get('单箱数量'), '单箱数量')
            if pack <= 0 or pack != int(pack):
                raise ValueError('单箱数量必须是正整数')
            stock = {h: number(inventory.get(h), h) for h in required_inv}
            available = stock['FBA可售']
            transfer = stock['FBA待调仓'] + stock['FBA调仓中']
            # Actual and marked quantities are alternative status snapshots, not additive.
            inbound = max(stock['FBA标发在途'], stock['FBA实际在途']) + stock['FBA入库中']
            awd_to_fba = max(stock['AWD补货至FBA标发在途'], stock['AWD补货至FBA实际在途'])
            # AWD->FBA may already appear in FBA transit: conservative union estimate.
            pipeline = max(inbound, awd_to_fba)
            awd_available = stock['AWD可用量']
            awd_inbound = max(stock['AWD标发在途'], stock['AWD实际在途'])
            fba_position = available + transfer + pipeline
            domestic, orders = 0.0, 0.0
            supply_valid = 'plan' in roles
            for role, row in plans[key]:
                row_shop = text(row.get('店铺'))
                if row_shop and row_shop != cfg['plan_store']:
                    supply_valid = False
                    issue(key, f'计划单第{row["__row__"]}行店铺为{row_shop}，请填写对应计划单店铺；未纳入')
                    continue
                quantity = number(row.get('总数量'), '计划单总数量')
                shipped = number(row.get('发货数量'), '计划单发货数量', optional=True)
                if shipped > quantity:
                    raise ValueError('计划单发货数量超过总数量')
                status = text(row.get('计划状态'))
                # User confirms completed quantities are current unshipped domestic stock.
                # Do not deduct historical shipment fields from that current balance again.
                remaining = quantity if status == '已完成' else quantity - shipped
                included = '未确认'
                if status == '已完成':
                    domestic += remaining
                    included = '国内可用现货（已确认）'
                elif status in ('已采购', '已排待做', '待入库'):
                    orders += remaining
                    included = '已有采购未交付'
                elif status in ('未下订单', '已取消', '取消'):
                    included = '不计入供给'
                elif remaining:
                    supply_valid = False
                    issue(key, f'计划单第{row["__row__"]}行状态“{status or "空白"}”待确认，采购与国内发货建议暂缓')
                source_rows.append({'FNSKU': key, '来源': roles[role], '行号': row['__row__'],
                                    '原数量': quantity, '已发数量': shipped, '剩余数量': remaining, '归类': included})
            coverage = (100 if demand > 49 else 90 if demand > 19 else 80) if cfg['coverage_mode'] == 'tiered' else cfg['target_days']
            shipping_horizon = cfg['shipping_days'] + coverage + cfg['safety_days']
            procurement_horizon = shipping_horizon + cfg['production_days']
            gap = max(0, demand * shipping_horizon - fba_position - awd_available)
            fba_boxes = math.ceil(gap / pack)
            purchase_gap = max(0, demand * procurement_horizon - fba_position
                               - awd_available - awd_inbound - domestic - orders)
            purchase_boxes = math.ceil(purchase_gap / pack)
            out.update({'采用日销': demand, '采用覆盖天数': coverage, '单箱数量': pack,
                        'FBA可售': available, 'FBA调仓供给': transfer,
                        'FBA在途估计': pipeline, 'AWD可用': awd_available,
                        'AWD在途估计': awd_inbound, '国内仓候选结余': domestic,
                        '已有采购候选未交付': orders, 'FBA可售天数': round(available / demand, 1) if demand else None,
                        'FBA补货需求箱数': fba_boxes, 'FBA补货需求量': fba_boxes * int(pack),
                        'AWD优先调拨参考量': min(awd_available, max(0, demand * shipping_horizon - fba_position)),
                        '采购缺口试算量': purchase_boxes * int(pack),
                        'FBA断货风险': '无销量，人工复核' if demand == 0 else
                        ('运输到货前可能断货' if available / demand < cfg['shipping_days'] else '未触发可售库存预警')})
            snapshot = catalog.get('备货日期')
            if isinstance(snapshot, datetime):
                out['原表备货日期'] = snapshot.date().isoformat()
                if snapshot.date() != cfg['as_of']:
                    issue(key, f'原表备货日期为{snapshot.date()}，与本次数据日期不同；请确认原始报表已更新')
            confirmed = cfg['data_confirmed'] and supply_valid
            if demand == 0:
                out['状态'] = '零销量待复核'
            elif confirmed:
                out['状态'] = '可复核建议'
                ship_boxes = min(fba_boxes, math.floor(domestic / pack))
                out.update({'建议FBA发货箱数': ship_boxes, '建议FBA发货量': ship_boxes * int(pack),
                            '国内现货不足量': max(0, fba_boxes * int(pack) - ship_boxes * int(pack)),
                            '建议采购箱数': purchase_boxes, '建议采购量': purchase_boxes * int(pack)})
            else:
                issue(key, '仅试算：请确认本次报表为同日快照及当前计划单结余')
        except ValueError as exc:
            out['状态'] = '数据错误'
            issue(key, str(exc))

    if not results:
        raise ValueError('商品主表未找到带 FNSKU 的商品行')
    shipment_rows = []
    for row in tables.get(roles.get('shipments'), []):
        key = text(row.get('条码(FNSKU或GTIN)'))
        if not key:
            continue
        entry = {'FNSKU': key, '来源行号': row['__row__'], '货件编号': '待补充',
                 '登记数量': row.get('总数量'), '接收数量': None, '差异': None,
                 '状态': '缺少后台接收数量，尚不能核对短收'}
        try:
            total = number(row.get('总数量'), '货件总数量')
            units = number(row.get('单箱数量'), '货件单箱数量')
            boxes = number(row.get('箱数'), '货件箱数')
            entry.update({'按箱计算数量': units * boxes, '装箱差异': total - units * boxes})
            if total != units * boxes:
                entry['状态'] = '登记数量与装箱数量不一致'
        except ValueError as exc:
            entry['状态'] = str(exc)
        shipment_rows.append(entry)

    assumptions = [
        '这是参数化试算模型，不是原表所有历史公式的复刻；不自动下单或创建货件。',
        'FBA补货需求=max(0,日销×(运输上架天数+目标覆盖天数+安全天数)−FBA供给−AWD可用)，向上取整箱。',
        '采购缺口=max(0,日销×(生产+运输上架+目标覆盖+安全天数)−FBA供给−AWD可用−AWD在途−国内结余−已有采购)，向上取整箱。',
        '生产45天、运输及上架50天；按原表销量梯度设置覆盖：日销>49为100天，>19为90天，其余80天，可切换手动天数。安全7天为可调整假设。',
        '建议FBA发货受国内已完成现货整箱库存限制；厂家已出/未出是袋子数据，完全排除本次补货及采购核算。',
        'FBA供给=可售+待调仓+调仓中+在途估计；不计预留、待发货、计划入库、不可售。',
        '标发/实际在途取较大值；FBA在途与AWD至FBA在途再取较大值，避免直接相加。此为保守估计，缺货件级唯一ID不能准确去重。',
        'AWD仅采用FBA库存明细中的AWD字段；不再叠加ADW库存表。AWD在途仅抵扣采购，不抵扣本次国内发货需求。',
        '用户已确认：计划单已完成的总数量就是国内仓尚未发往FBA的现货，不再减历史发货字段；已采购/已排待做/待入库取总数量减发货数量作为已有订单。',
        '空白/缺料等未知计划状态会阻止该商品的正式建议；须确认本次各表是同一日期的当前结余，非历史累计。',
        '同一店铺+FNSKU重复时报错；缺记录、空销量、非法数字不按零处理；主表缺FNSKU的空白占位行跳过。',
        '零销量不自动给出采购建议；在途未按预计到货日分段，时间风险须人工复核。',
        '已创数据只有登记及装箱信息，不能据此确定实际接收数量；不将全部历史已创货件再次加入在途。',
    ]
    return {'results': results, 'issues': issues, 'shipments': shipment_rows,
            'supply_details': source_rows, 'assumptions': assumptions,
            'settings': {**cfg, 'as_of': cfg['as_of'].isoformat(), 'shop': shop},
            'sources': [{'工作表': k, '读取非空行数': len(v)} for k, v in tables.items()]}


def export_workbook(report):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    datasets = [('补货与采购', report['results']), ('数据待办', report['issues']),
                ('货件装箱核对', report['shipments']), ('供给扣减明细', report['supply_details']),
                ('来源工作表', report['sources']),
                ('参数与口径', [{'项目': k, '值': v} for k, v in report['settings'].items()] +
                 [{'项目': '计算说明', '值': x} for x in report['assumptions']])]
    for name, rows in datasets:
        sheet = wb.create_sheet(name)
        heads = list(dict.fromkeys(k for row in rows for k in row)) or ['说明']
        sheet.append(heads)
        for row in rows:
            values = []
            for head in heads:
                value = row.get(head)
                # All imported strings are text, including formula-like SKU/name values.
                if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                    value = "'" + value
                values.append(value)
            sheet.append(values)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='183C38')
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(60, max(16, max(len(str(c.value or '')) for c in col[:40]) + 3))
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
